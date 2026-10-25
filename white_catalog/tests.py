import base64
import io
import json
import re
import tempfile

from decimal import Decimal

from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.db import IntegrityError
from django.test import TestCase, override_settings
from django.urls import reverse

from .models import (
    WhiteCart,
    WhiteCartItem,
    WhiteCatalogUser,
    WhiteCategory,
    WhiteColor,
    WhiteColorVariant,
    WhiteFabricType,
    WhiteOrder,
    WhiteOrderItem,
    WhitePackType,
    WhiteProductVariant,
    WhiteSizeType,
    WhiteSubcategory,
    WhiteSubcategoryImage,
    WhiteSubcategoryPrice,
    apply_price_list,
)
from .templatetags.white_catalog_extras import has_price_discount
from .views import _product_gallery_images


class WhiteCatalogAdminBridgeTests(TestCase):
    def setUp(self):
        self.django_admin = get_user_model().objects.create_user(
            username="admin",
            password="lior1234",
            is_staff=True,
            is_superuser=True,
        )

    def test_white_catalog_login_accepts_django_admin_credentials(self):
        response = self.client.post(
            reverse("white_catalog:login"),
            {
                "username": "admin",
                "password": "lior1234",
                "next": reverse("white_catalog:cart"),
            },
        )

        self.assertRedirects(response, reverse("white_catalog:cart"))
        catalog_user = WhiteCatalogUser.objects.get(username="admin")
        self.assertEqual(self.client.session["white_catalog_user_id"], catalog_user.id)
        self.assertTrue(catalog_user.check_password("lior1234"))
        self.assertTrue(catalog_user.is_active)

    def test_logged_in_django_admin_can_open_catalog_without_separate_login(self):
        self.client.force_login(self.django_admin)

        response = self.client.get(reverse("white_catalog:cart"))

        self.assertEqual(response.status_code, 200)
        self.assertTrue(WhiteCatalogUser.objects.filter(username="admin", is_active=True).exists())

    def test_staff_catalog_logout_clears_session_and_shows_login(self):
        self.client.force_login(self.django_admin)
        self.client.get(reverse("white_catalog:home"))
        self.assertTrue(self.client.session.get("white_catalog_user_id"))

        response = self.client.get(reverse("white_catalog:logout"), follow=True)

        self.assertIsNone(self.client.session.get("white_catalog_user_id"))
        self.assertContains(response, reverse("white_catalog:login"))
        self.assertContains(response, "התחבר")
        self.assertNotContains(response, reverse("white_catalog:logout"))
        self.assertNotContains(response, "Admin Test Account")


@override_settings(WHITE_CATALOG_API_TOKEN="test-token", MEDIA_ROOT=tempfile.mkdtemp())
class ImportColorVariantsTests(TestCase):
    """ייבוא שורות צבע מהתוכנה השולחנית אל 'צבעים למוצר'."""

    def setUp(self):
        self.product = WhiteSubcategory.objects.create(name="סדין למיטת תינוק")
        self.url = reverse("white_catalog:api_import_variants")

    def _post(self, rows, fabric="טריקו ג'רזי"):
        return self.client.post(
            self.url,
            data=json.dumps({"product_id": self.product.id, "fabric_type": fabric, "rows": rows}),
            content_type="application/json",
            headers={"X-API-Token": "test-token"},
        )

    @staticmethod
    def _jpeg_b64():
        from PIL import Image

        buf = io.BytesIO()
        Image.new("RGB", (600, 600), (115, 135, 168)).save(buf, "JPEG")
        return base64.b64encode(buf.getvalue()).decode("ascii")

    def test_color_rows_create_color_variants_without_size(self):
        rows = [
            {"barcode": "7297555022264", "unit_price": "21.3", "color": "תכלת קשי", "color_hex": "#7387a8"},
            {"barcode": "7297555022271", "unit_price": "21.3", "color": "ורוד קשי", "color_hex": "#FFA4A4"},
        ]
        response = self._post(rows)

        data = response.json()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(data["errors"], [])
        self.assertEqual(data["created"], 2)
        self.assertEqual(WhiteColorVariant.objects.filter(product=self.product).count(), 2)
        tchelet = WhiteColorVariant.objects.get(barcode="7297555022264")
        self.assertEqual(tchelet.color.name, "תכלת קשי")
        self.assertEqual(tchelet.color.hex_color, "#7387a8")
        self.assertIsNone(tchelet.size_id)
        self.assertEqual(str(tchelet.unit_price), "21.30")
        # המוצר סומן אוטומטית להזמנה לפי מניפת צבעים
        self.product.refresh_from_db()
        self.assertTrue(self.product.has_color_variants)
        self.assertTrue(self.product.is_orderable)
        # לא נוצרו וריאנטים לפי מידה מהשורות האלה
        self.assertEqual(WhiteProductVariant.objects.count(), 0)

    def test_color_rows_are_idempotent_and_update_price(self):
        self._post([{"barcode": "111", "unit_price": "20", "color": "תכלת"}])
        response = self._post([{"barcode": "111", "unit_price": "25", "color": "תכלת"}])

        data = response.json()
        self.assertEqual(data["created"], 0)
        self.assertEqual(data["updated"], 1)
        variant = WhiteColorVariant.objects.get(barcode="111")
        self.assertEqual(str(variant.unit_price), "25.00")
        self.assertEqual(WhiteColorVariant.objects.count(), 1)
        self.assertEqual(WhiteColor.objects.count(), 1)

    def test_color_row_saves_variant_image_and_color_swatch(self):
        rows = [{
            "barcode": "222",
            "unit_price": "21.3",
            "color": "תכלת",
            "color_hex": "#7387a8",
            "image_base64": self._jpeg_b64(),
            "image_format": "jpg",
        }]
        response = self._post(rows)

        self.assertEqual(response.json()["errors"], [])
        variant = WhiteColorVariant.objects.get(barcode="222")
        self.assertTrue(variant.image)
        self.assertTrue(variant.image.name.endswith(".jpg"))
        # התמונה נשמרת גם כדוגמית הצבע במניפה
        color = WhiteColor.objects.get(name="תכלת")
        self.assertTrue(color.swatch_image)
        self.assertTrue(color.swatch_image.name.endswith(".jpg"))
        self.product.refresh_from_db()
        self.assertTrue(self.product.image)
        self.assertTrue(self.product.image.name.endswith(".jpg"))

    def test_color_import_switches_product_from_order_variants_to_color_mode(self):
        self.product.has_order_variants = True
        self.product.save(update_fields=["has_order_variants"])

        response = self._post([{"barcode": "666", "unit_price": "21.3", "color": "תכלת"}])

        data = response.json()
        self.assertEqual(data["errors"], [])
        self.product.refresh_from_db()
        self.assertTrue(self.product.has_color_variants)
        self.assertFalse(self.product.has_order_variants)
        self.assertTrue(any("כובה" in w for w in data["warnings"]))

    def test_color_row_releases_barcode_from_stale_size_variant(self):
        fabric = WhiteFabricType.objects.create(name="טריקו ג'רזי")
        size = WhiteSizeType.objects.create(name="66x128")
        WhiteProductVariant.objects.create(
            product=self.product, fabric_type=fabric, size_type=size, barcode="333"
        )

        response = self._post([{"barcode": "333", "unit_price": "21.3", "color": "שמנת"}])

        self.assertEqual(response.json()["errors"], [])
        self.assertTrue(WhiteColorVariant.objects.filter(barcode="333").exists())
        stale = WhiteProductVariant.objects.get(product=self.product)
        self.assertIsNone(stale.barcode)

    def test_color_row_conflicting_barcode_on_other_product_is_rejected(self):
        other = WhiteSubcategory.objects.create(name="מוצר אחר")
        color = WhiteColor.objects.create(name="לבן")
        WhiteColorVariant.objects.create(product=other, color=color, barcode="444")

        response = self._post([{"barcode": "444", "unit_price": "10", "color": "לבן"}])

        data = response.json()
        self.assertEqual(data["created"], 0)
        self.assertEqual(len(data["errors"]), 1)
        self.assertIn("444", data["errors"][0])

    def test_clear_image_keeps_hex_only_display(self):
        rows = [{
            "barcode": "222",
            "unit_price": "21.3",
            "color": "תכלת",
            "color_hex": "#7387a8",
            "image_base64": self._jpeg_b64(),
            "image_format": "jpg",
        }]
        self._post(rows)
        variant = WhiteColorVariant.objects.get(barcode="222")
        self.assertTrue(variant.image)
        self.assertTrue(variant.color.swatch_image)

        response = self._post([{
            "barcode": "222",
            "unit_price": "21.3",
            "color": "תכלת",
            "color_hex": "#7387a8",
            "clear_image": True,
        }])
        self.assertEqual(response.json()["errors"], [])
        variant.refresh_from_db()
        color = WhiteColor.objects.get(name="תכלת")
        self.assertFalse(variant.image)
        self.assertFalse(color.swatch_image)
        self.assertEqual(color.hex_color, "#7387a8")

    def test_delete_color_variant_by_barcode(self):
        color = WhiteColor.objects.create(name="אפור עשן בהיר")
        WhiteColorVariant.objects.create(product=self.product, color=color, barcode="7297555022295")

        response = self._post([{"barcode": "7297555022295", "delete": True}])

        data = response.json()
        self.assertEqual(data["errors"], [])
        self.assertEqual(data["deleted"], 1)
        self.assertFalse(WhiteColorVariant.objects.filter(barcode="7297555022295").exists())

    def test_size_rows_still_work_as_before(self):
        rows = [
            {"size": "0-3", "barcode": "555", "unit_price": "12.5"},
            {"size": "3-6", "barcode": "556", "unit_price": "12.5"},
        ]
        response = self._post(rows)

        data = response.json()
        self.assertEqual(data["errors"], [])
        self.assertEqual(data["created"], 2)
        self.assertEqual(WhiteProductVariant.objects.count(), 2)
        self.product.refresh_from_db()
        self.assertTrue(self.product.has_order_variants)
        self.assertFalse(self.product.has_color_variants)


class ClothingThreesOnlyTests(TestCase):
    """ביגוד מוזמן בשלישיות בלבד — חמישיות לא מוצעות ולא מתקבלות."""

    def setUp(self):
        self.threes = WhitePackType.objects.create(name="שלישיות", quantity=3, is_active=True)
        self.fives = WhitePackType.objects.create(name="חמישיות", quantity=5, is_active=True)
        self.user = WhiteCatalogUser.objects.create(
            company_name="בדיקה",
            contact_name="בודק",
            contact_phone="050",
            username="tester",
            pack_route=WhiteCatalogUser.ROUTE_FIVES,
        )
        self.user.set_password("pass")
        self.user.save()
        session = self.client.session
        session["white_catalog_user_id"] = self.user.id
        session["white_catalog_username"] = self.user.username
        session.save()

        self.product = WhiteSubcategory.objects.create(
            name="בגד גוף",
            is_orderable=True,
            has_order_variants=True,
        )
        fabric = WhiteFabricType.objects.create(name="טריקו")
        size, _created = WhiteSizeType.objects.get_or_create(name="0-3")
        self.variant = WhiteProductVariant.objects.create(
            product=self.product,
            fabric_type=fabric,
            size_type=size,
            unit_price="12.00",
            barcode="111",
        )
        self.variant.pack_types.set([self.threes, self.fives])

    def test_allowed_pack_types_are_threes_even_if_user_route_is_fives(self):
        ids = list(self.user.get_allowed_pack_types().values_list("id", flat=True))
        self.assertEqual(ids, [self.threes.id])

    def test_cart_add_rejects_fives(self):
        response = self.client.post(
            reverse("white_catalog:cart_add"),
            {
                "format": "json",
                "pack_type_id": str(self.fives.id),
                f"qty_{self.variant.id}": "2",
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("שלישיות", response.json()["error"])

    def test_cart_add_accepts_threes(self):
        response = self.client.post(
            reverse("white_catalog:cart_add"),
            {
                "format": "json",
                "pack_type_id": str(self.threes.id),
                f"qty_{self.variant.id}": "2",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["ok"])


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class MergePrintedTetraPagesTests(TestCase):
    def setUp(self):
        self.category = WhiteCategory.objects.create(
            name="חיתולי בד / טטרה",
            slug="baby_diapers",
        )
        self.keep = WhiteSubcategory.objects.create(
            name="מארז 3 חיתולי טטרה - מודפס",
            slug="muslin_baby_diapers",
            category=self.category,
            is_orderable=True,
        )

    def test_legacy_hebrew_slug_redirects_to_muslin_page(self):
        response = self.client.get(
            reverse(
                "white_catalog:subcategory_detail",
                kwargs={
                    "category_slug": "baby_diapers",
                    "subcategory_slug": "מארז-חיתולי-טטרה-מודפס-3",
                },
            )
        )
        self.assertEqual(response.status_code, 301)
        self.assertEqual(
            response.url,
            reverse(
                "white_catalog:subcategory_detail",
                kwargs={
                    "category_slug": "baby_diapers",
                    "subcategory_slug": "muslin_baby_diapers",
                },
            ),
        )

    def test_gallery_includes_product_and_print_images_without_duplicates(self):
        self.keep.image.save("main.jpg", ContentFile(b"main"), save=True)
        color = WhiteColor.objects.create(name="ינשופים כחול")
        variant = WhiteColorVariant.objects.create(product=self.keep, color=color)
        variant.image.save("print.jpg", ContentFile(b"print"), save=True)
        gallery = _product_gallery_images(self.keep)
        urls = [item["url"] for item in gallery]
        self.assertEqual(len(urls), 2)
        self.assertEqual(len(set(urls)), 2)
        self.assertTrue(any("main" in url for url in urls))
        self.assertTrue(any("print" in url for url in urls))

    def test_gallery_hides_same_photo_stored_as_different_files(self):
        from PIL import Image

        buf = io.BytesIO()
        Image.new("RGB", (80, 80), (200, 160, 140)).save(buf, "JPEG", quality=90)
        photo = buf.getvalue()
        extra = WhiteSubcategoryImage.objects.create(subcategory=self.keep, order=1)
        extra.image.save("gallery.jpg", ContentFile(photo), save=True)
        color = WhiteColor.objects.create(name="אוהלים וכוכבים ורוד")
        variant = WhiteColorVariant.objects.create(product=self.keep, color=color)
        variant.image.save("variant.jpg", ContentFile(photo), save=True)
        gallery = _product_gallery_images(self.keep)
        self.assertEqual(len(gallery), 1)

    def test_gallery_hides_reencoded_copy_of_the_same_photo(self):
        from PIL import Image

        src = Image.new("RGB", (120, 80), (40, 90, 140))
        src.paste((220, 180, 160), (20, 10, 70, 60))
        high = io.BytesIO()
        low = io.BytesIO()
        src.save(high, "JPEG", quality=92)
        src.resize((1400, 933)).save(low, "JPEG", quality=88)
        extra = WhiteSubcategoryImage.objects.create(subcategory=self.keep, order=1)
        extra.image.save("orig.jpg", ContentFile(high.getvalue()), save=True)
        color = WhiteColor.objects.create(name="סוסים וזיגזגים כחול")
        variant = WhiteColorVariant.objects.create(product=self.keep, color=color)
        variant.image.save("export.jpg", ContentFile(low.getvalue()), save=True)
        gallery = _product_gallery_images(self.keep)
        self.assertEqual(len(gallery), 1)


class ApplyPriceListTests(TestCase):
    def test_catalog_keeps_list_price(self):
        user = WhiteCatalogUser(
            company_name="א",
            contact_name="ב",
            contact_phone="1",
            username="c1",
            price_percent=Decimal("0.00"),
        )
        self.assertEqual(apply_price_list(Decimal("13.00"), user), Decimal("13.00"))

    def test_ten_percent_discount_and_bodysuit_tiers(self):
        user_10 = WhiteCatalogUser(
            company_name="א",
            contact_name="ב",
            contact_phone="1",
            username="d10",
            price_percent=Decimal("10.00"),
        )
        user_l = WhiteCatalogUser(
            company_name="א",
            contact_name="ב",
            contact_phone="1",
            username="l",
            price_percent=Decimal("7.69"),
        )
        self.assertEqual(apply_price_list(Decimal("13.00"), user_10), Decimal("11.7"))
        self.assertEqual(apply_price_list(Decimal("13.00"), user_l), Decimal("12.0"))

    def test_discount_rounds_to_one_decimal(self):
        user_15 = WhiteCatalogUser(
            company_name="א",
            contact_name="ב",
            contact_phone="1",
            username="d15",
            price_percent=Decimal("15.00"),
        )
        self.assertEqual(apply_price_list(Decimal("12.80"), user_15), Decimal("10.9"))
        self.assertEqual(apply_price_list(Decimal("13.60"), user_15), Decimal("11.6"))
        self.assertEqual(str(apply_price_list(Decimal("12.80"), user_15)), "10.9")
        self.assertEqual(str(apply_price_list(Decimal("13.60"), user_15)), "11.6")
        self.assertEqual(apply_price_list(Decimal("12.80"), None), Decimal("12.80"))

    def test_none_price_stays_none(self):
        self.assertIsNone(apply_price_list(None, None))

    def test_has_price_discount_filter(self):
        user_0 = WhiteCatalogUser(price_percent=Decimal("0.00"))
        user_10 = WhiteCatalogUser(price_percent=Decimal("10.00"))
        self.assertFalse(has_price_discount(None))
        self.assertFalse(has_price_discount(user_0))
        self.assertTrue(has_price_discount(user_10))


class CustomerPriceListTests(TestCase):
    def setUp(self):
        self.threes = WhitePackType.objects.create(name="שלישיות", quantity=3, is_active=True)
        self.product = WhiteSubcategory.objects.create(
            name="בגד גוף שרוול ארוך",
            slug="Baby_long_bodysuit_test",
            is_orderable=True,
            has_order_variants=True,
            unit_price=Decimal("13.00"),
        )
        fabric = WhiteFabricType.objects.create(name="פלנל")
        size, _created = WhiteSizeType.objects.get_or_create(name="0-3")
        self.variant = WhiteProductVariant.objects.create(
            product=self.product,
            fabric_type=fabric,
            size_type=size,
            unit_price=Decimal("13.00"),
            barcode="7297555020987",
        )
        self.variant.pack_types.set([self.threes])

    def _login(self, username, price_percent):
        user = WhiteCatalogUser.objects.create(
            company_name="בדיקה",
            contact_name="בודק",
            contact_phone="050",
            username=username,
            price_percent=price_percent,
        )
        user.set_password("pass")
        user.save()
        session = self.client.session
        session["white_catalog_user_id"] = user.id
        session["white_catalog_username"] = user.username
        session.save()
        return user

    def test_catalog_customer_sees_13_and_cart_pack_is_39(self):
        user = self._login("shop", Decimal("0.00"))
        url = reverse(
            "white_catalog:standalone_subcategory_detail",
            kwargs={"subcategory_slug": self.product.slug},
        )
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        sizes = response.context["variants_data"][0]["sizes"]
        self.assertEqual(sizes[0]["unit_price"], "13.00")
        self.assertEqual(sizes[0]["list_unit_price"], "13.00")
        self.assertNotContains(response, "המחירים המוצגים הם אחרי ההנחה")
        self.assertNotContains(response, "יש לך")

        add = self.client.post(
            reverse("white_catalog:cart_add"),
            {
                "format": "json",
                "pack_type_id": str(self.threes.id),
                f"qty_{self.variant.id}": "1",
            },
        )
        self.assertEqual(add.status_code, 200)
        item = WhiteCartItem.objects.get(cart__user=user, variant=self.variant)
        self.assertEqual(item.price_at_add, Decimal("39.00"))

    def test_ten_percent_customer_sees_11_70_and_cart_pack_is_35_10(self):
        user = self._login("disc10", Decimal("10.00"))
        url = reverse(
            "white_catalog:standalone_subcategory_detail",
            kwargs={"subcategory_slug": self.product.slug},
        )
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        sizes = response.context["variants_data"][0]["sizes"]
        self.assertEqual(sizes[0]["unit_price"], "11.7")
        self.assertEqual(sizes[0]["list_unit_price"], "13.00")
        self.assertContains(response, "יש לך 10% הנחה")
        self.assertContains(response, "המחירים המוצגים הם אחרי ההנחה")

        add = self.client.post(
            reverse("white_catalog:cart_add"),
            {
                "format": "json",
                "pack_type_id": str(self.threes.id),
                f"qty_{self.variant.id}": "2",
            },
        )
        self.assertEqual(add.status_code, 200)
        item = WhiteCartItem.objects.get(cart__user=user, variant=self.variant)
        self.assertEqual(item.price_at_add, Decimal("35.10"))
        self.assertEqual(item.quantity, 2)


class PerUserCatalogVisibilityTests(TestCase):
    def setUp(self):
        self.category = WhiteCategory.objects.create(name="ביגוד", slug="clothing")
        self.other_category = WhiteCategory.objects.create(name="מצעים", slug="bedding")
        self.hidden_product = WhiteSubcategory.objects.create(
            name="בגד גוף מוסתר",
            slug="hidden-bodysuit",
            category=self.category,
            is_orderable=True,
            unit_price=Decimal("13.00"),
        )
        self.visible_product = WhiteSubcategory.objects.create(
            name="בגד גוף גלוי",
            slug="visible-bodysuit",
            category=self.category,
            is_orderable=True,
            unit_price=Decimal("13.00"),
        )
        self.other_product = WhiteSubcategory.objects.create(
            name="סדין",
            slug="sheet",
            category=self.other_category,
            is_orderable=True,
            unit_price=Decimal("20.00"),
        )

    def _login(self, username="shop-hide"):
        user = WhiteCatalogUser.objects.create(
            company_name="חנות בדיקה",
            contact_name="בודק",
            contact_phone="050",
            username=username,
        )
        user.set_password("pass")
        user.save()
        session = self.client.session
        session["white_catalog_user_id"] = user.id
        session["white_catalog_username"] = user.username
        session.save()
        return user

    def _category_url(self, category):
        return reverse("white_catalog:category_detail", kwargs={"category_slug": category.slug})

    def _product_url(self, product):
        return reverse(
            "white_catalog:subcategory_detail",
            kwargs={
                "category_slug": product.category.slug,
                "subcategory_slug": product.slug,
            },
        )

    def test_guest_sees_full_catalog(self):
        home = self.client.get(reverse("white_catalog:home"))
        self.assertEqual(home.status_code, 200)
        self.assertContains(home, self.category.name)
        self.assertContains(home, self.other_category.name)

        category_page = self.client.get(self._category_url(self.category))
        self.assertEqual(category_page.status_code, 200)
        self.assertContains(category_page, self.hidden_product.name)
        self.assertContains(category_page, self.visible_product.name)
        self.assertEqual(self.client.get(self._product_url(self.hidden_product)).status_code, 200)

    def test_user_without_hides_sees_full_catalog(self):
        self._login()
        home = self.client.get(reverse("white_catalog:home"))
        self.assertEqual(home.status_code, 200)
        nav_ids = list(home.context["all_categories"].values_list("id", flat=True))
        self.assertIn(self.category.id, nav_ids)
        self.assertIn(self.other_category.id, nav_ids)
        self.assertEqual(self.client.get(self._category_url(self.category)).status_code, 200)
        self.assertEqual(self.client.get(self._product_url(self.visible_product)).status_code, 200)

    def test_hidden_category_is_gone_for_that_user(self):
        user = self._login()
        user.hidden_categories.add(self.category)

        home = self.client.get(reverse("white_catalog:home"))
        nav_ids = list(home.context["all_categories"].values_list("id", flat=True))
        self.assertNotIn(self.category.id, nav_ids)
        self.assertIn(self.other_category.id, nav_ids)
        self.assertEqual(self.client.get(self._category_url(self.category)).status_code, 404)
        self.assertEqual(self.client.get(self._product_url(self.visible_product)).status_code, 404)
        self.assertEqual(self.client.get(self._product_url(self.other_product)).status_code, 200)

    def test_hidden_product_is_gone_and_cannot_be_added_to_cart(self):
        user = self._login()
        user.hidden_products.add(self.hidden_product)

        category_page = self.client.get(self._category_url(self.category))
        self.assertEqual(category_page.status_code, 200)
        names = [product.name for product in category_page.context["subcategories"]]
        self.assertNotIn(self.hidden_product.name, names)
        self.assertIn(self.visible_product.name, names)
        self.assertEqual(self.client.get(self._product_url(self.hidden_product)).status_code, 404)
        self.assertEqual(self.client.get(self._product_url(self.visible_product)).status_code, 200)

        add = self.client.post(
            reverse("white_catalog:cart_add"),
            {
                "format": "json",
                "product_id": str(self.hidden_product.id),
                "simple_quantity": "1",
            },
        )
        self.assertEqual(add.status_code, 400)
        self.assertFalse(
            WhiteCartItem.objects.filter(cart__user=user, product=self.hidden_product).exists()
        )


class StoreImportExportTests(TestCase):
    def setUp(self):
        self.threes = WhitePackType.objects.create(name="שלישיות", quantity=3, is_active=True)
        self.user = WhiteCatalogUser.objects.create(
            company_name="Toyland",
            contact_name="איילת",
            contact_phone="050",
            username="toyland_test",
        )
        self.user.set_password("pass")
        self.user.save()
        session = self.client.session
        session["white_catalog_user_id"] = self.user.id
        session["white_catalog_username"] = self.user.username
        session.save()

        self.category = WhiteCategory.objects.create(name="לבנים", slug="basics-test")
        fabric = WhiteFabricType.objects.create(name="פלנל")
        size_a, _created = WhiteSizeType.objects.get_or_create(name="0-3")
        size_b = WhiteSizeType.objects.create(name="3-6")
        self.product = WhiteSubcategory.objects.create(
            name="בגד גוף חורף",
            slug="winter-bodysuit",
            category=self.category,
            is_orderable=True,
            has_order_variants=True,
            unit_price=Decimal("13.00"),
            online_price=Decimal("29.90"),
            marketing_description="לבנים לחורף",
        )
        self.v1 = WhiteProductVariant.objects.create(
            product=self.product,
            fabric_type=fabric,
            size_type=size_a,
            unit_price=Decimal("13.00"),
            barcode="1001",
        )
        self.v2 = WhiteProductVariant.objects.create(
            product=self.product,
            fabric_type=fabric,
            size_type=size_b,
            unit_price=Decimal("13.00"),
            barcode="1002",
        )
        self.hidden = WhiteSubcategory.objects.create(
            name="מוצר מוסתר",
            slug="hidden-item",
            is_orderable=True,
            unit_price=Decimal("5.00"),
        )
        self.user.hidden_products.add(self.hidden)

    def _csv_rows(self, response):
        import csv
        from io import StringIO

        text = response.content.decode("utf-8-sig")
        return list(csv.DictReader(StringIO(text)))

    def test_store_csv_groups_sizes_under_one_handle(self):
        response = self.client.get(reverse("white_catalog:export_store_import"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response["Content-Type"])
        rows = self._csv_rows(response)
        variant_rows = [r for r in rows if r["Variant Barcode"]]
        self.assertEqual({r["Handle"] for r in variant_rows}, {"winter-bodysuit"})
        self.assertEqual({r["Variant Barcode"] for r in variant_rows}, {"1001", "1002"})
        self.assertEqual(variant_rows[0]["Option1 Name"], "מידה")
        self.assertEqual(variant_rows[0]["Option2 Name"], "בד")
        self.assertEqual(variant_rows[0]["Variant Price"], "29.90")
        self.assertEqual(variant_rows[0]["Cost per item"], "13.00")
        self.assertNotIn("hidden-item", {r["Handle"] for r in rows})

    def test_product_scope_csv_and_excel(self):
        extra = WhiteSubcategory.objects.create(
            name="מוצר נוסף",
            slug="other-item",
            is_orderable=True,
            unit_price=Decimal("8.00"),
        )
        csv_response = self.client.get(
            reverse("white_catalog:export_store_import"),
            {"product": "winter-bodysuit"},
        )
        self.assertEqual(csv_response.status_code, 200)
        handles = {r["Handle"] for r in self._csv_rows(csv_response)}
        self.assertEqual(handles, {"winter-bodysuit"})
        self.assertNotIn(extra.slug, handles)

        excel_response = self.client.get(
            reverse("white_catalog:export_products"),
            {"product": "winter-bodysuit"},
        )
        self.assertEqual(excel_response.status_code, 200)
        self.assertIn("winter-bodysuit", excel_response["Content-Disposition"])
        from openpyxl import load_workbook
        from io import BytesIO
        wb = load_workbook(BytesIO(excel_response.content))
        barcodes = [row[0].value for row in wb.active.iter_rows(min_row=2, max_col=1)]
        self.assertEqual(set(barcodes), {"1001", "1002"})

    def test_order_scope_csv_only_ordered_size(self):
        order = WhiteOrder.objects.create(
            user=self.user,
            order_number="ORD-TEST-1",
            total_amount=Decimal("39.00"),
        )
        WhiteOrderItem.objects.create(
            order=order,
            product=self.product,
            variant=self.v1,
            product_name=self.product.name,
            variant_name="פלנל",
            barcode="1001",
            size_name="0-3",
            pack_type_name="שלישיות",
            pack_quantity=3,
            quantity=1,
            unit_price=Decimal("39.00"),
        )
        response = self.client.get(
            reverse("white_catalog:export_store_import"),
            {"order": "ORD-TEST-1"},
        )
        self.assertEqual(response.status_code, 200)
        barcodes = {r["Variant Barcode"] for r in self._csv_rows(response) if r["Variant Barcode"]}
        self.assertEqual(barcodes, {"1001"})

    def test_unknown_product_export_is_404(self):
        response = self.client.get(
            reverse("white_catalog:export_store_import"),
            {"product": "does-not-exist"},
        )
        self.assertEqual(response.status_code, 404)

    def _attach_jpegs(self, product, count=2):
        from PIL import Image

        for index in range(count):
            buf = io.BytesIO()
            Image.new("RGB", (160, 100), (30 * index, 90, 150)).save(buf, "JPEG")
            buf.seek(0)
            content = ContentFile(buf.getvalue(), name=f"export-{index}.jpg")
            if index == 0:
                product.image.save(f"export-{index}.jpg", content, save=True)
            else:
                WhiteSubcategoryImage.objects.create(
                    subcategory=product,
                    image=content,
                    alt_text=f"gallery-{index}",
                )

    @override_settings(MEDIA_ROOT=tempfile.mkdtemp())
    def test_catalog_excel_splits_images_into_columns(self):
        self._attach_jpegs(self.product, 2)
        response = self.client.get(
            reverse("white_catalog:export_products"),
            {"product": "winter-bodysuit"},
        )
        self.assertEqual(response.status_code, 200)
        from openpyxl import load_workbook
        from io import BytesIO
        wb = load_workbook(BytesIO(response.content))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        self.assertIn("תמונה 1", headers)
        self.assertIn("תמונה 2", headers)
        self.assertNotIn("קישורי תמונות", headers)
        img1_col = headers.index("תמונה 1") + 1
        img2_col = headers.index("תמונה 2") + 1
        first = ws.cell(2, img1_col).value
        second = ws.cell(2, img2_col).value
        self.assertTrue(first)
        self.assertTrue(second)
        self.assertNotIn("\n", first)
        self.assertNotIn("\n", second)
        self.assertEqual(len(ws._images), 0)

    @override_settings(MEDIA_ROOT=tempfile.mkdtemp())
    def test_order_excel_splits_image_columns(self):
        self._attach_jpegs(self.product, 2)
        order = WhiteOrder.objects.create(
            user=self.user,
            order_number="ORD-TEST-IMG",
            total_amount=Decimal("39.00"),
        )
        WhiteOrderItem.objects.create(
            order=order,
            product=self.product,
            variant=self.v1,
            product_name=self.product.name,
            variant_name="פלנל",
            barcode="1001",
            size_name="0-3",
            pack_type_name="שלישיות",
            pack_quantity=3,
            quantity=1,
            unit_price=Decimal("39.00"),
        )
        WhiteOrderItem.objects.create(
            order=order,
            product=self.product,
            variant=self.v2,
            product_name=self.product.name,
            variant_name="פלנל",
            barcode="1002",
            size_name="3-6",
            pack_type_name="שלישיות",
            pack_quantity=3,
            quantity=1,
            unit_price=Decimal("39.00"),
        )
        response = self.client.get(
            reverse("white_catalog:export_order", args=["ORD-TEST-IMG"]),
        )
        self.assertEqual(response.status_code, 200)
        from openpyxl import load_workbook
        from io import BytesIO
        wb = load_workbook(BytesIO(response.content))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        self.assertIn("תמונה 1", headers)
        self.assertNotIn("קישורי תמונות", headers)
        img1 = ws.cell(2, headers.index("תמונה 1") + 1).value
        self.assertTrue(img1)
        self.assertNotIn("\n", img1)
        self.assertEqual(len(ws._images), 0)


class CheckoutResetsCartTests(TestCase):
    def setUp(self):
        self.threes = WhitePackType.objects.create(name="שלישיות", quantity=3, is_active=True)
        self.user = WhiteCatalogUser.objects.create(
            company_name="בדיקה",
            contact_name="בודק",
            contact_phone="050",
            username="checkout-tester",
            pack_route=WhiteCatalogUser.ROUTE_THREES,
        )
        self.user.set_password("pass")
        self.user.save()
        session = self.client.session
        session["white_catalog_user_id"] = self.user.id
        session["white_catalog_username"] = self.user.username
        session.save()

        self.category = WhiteCategory.objects.create(name="ביגוד", slug="clothing-reset")
        self.product = WhiteSubcategory.objects.create(
            name="בגד גוף לחורף",
            slug="winter-bodysuit-reset",
            category=self.category,
            is_orderable=True,
            has_order_variants=True,
            unit_price=Decimal("13.00"),
        )
        fabric = WhiteFabricType.objects.create(name="פלנל-reset")
        size = WhiteSizeType.objects.create(name="0-3-reset")
        self.variant = WhiteProductVariant.objects.create(
            product=self.product,
            fabric_type=fabric,
            size_type=size,
            unit_price=Decimal("13.00"),
            barcode="2001",
        )
        self.variant.pack_types.set([self.threes])

    def _product_url(self):
        return reverse(
            "white_catalog:subcategory_detail",
            args=[self.category.slug, self.product.slug],
        )

    def _variants_data(self, response):
        match = re.search(
            r'<script id="variants-data" type="application/json">(.*?)</script>',
            response.content.decode(),
            re.S,
        )
        self.assertIsNotNone(match)
        return json.loads(match.group(1))

    def test_checkout_clears_product_quantities_and_rejects_stale_cart_add(self):
        add_response = self.client.post(
            reverse("white_catalog:cart_add"),
            {
                "format": "json",
                "pack_type_id": str(self.threes.id),
                f"qty_{self.variant.id}": "2",
            },
        )
        self.assertEqual(add_response.status_code, 200)
        self.assertTrue(add_response.json()["ok"])

        before = self._variants_data(self.client.get(self._product_url()))
        self.assertEqual(before[0]["sizes"][0]["cart_quantities"][str(self.threes.id)], 2)

        checkout = self.client.post(reverse("white_catalog:checkout"))
        self.assertEqual(checkout.status_code, 302)
        confirm = self.client.get(checkout.url)
        self.assertEqual(confirm.status_code, 200)
        self.assertContains(confirm, "window.__WHITE_CATALOG_RESET_CART = true")
        self.assertIn("white_catalog_cart_reset_at", checkout.cookies)
        self.assertTrue(WhiteOrder.objects.filter(user=self.user).exists())
        self.assertFalse(
            WhiteCart.objects.filter(user=self.user, status=WhiteCart.STATUS_ACTIVE).exists()
        )

        after = self._variants_data(self.client.get(self._product_url()))
        self.assertEqual(after[0]["sizes"][0]["cart_quantities"][str(self.threes.id)], 0)
        self.assertNotContains(self.client.get(checkout.url), "window.__WHITE_CATALOG_RESET_CART = true")

        stale = self.client.post(
            reverse("white_catalog:cart_add"),
            {
                "format": "json",
                "cart_generation": "0",
                "pack_type_id": str(self.threes.id),
                f"qty_{self.variant.id}": "4",
            },
        )
        self.assertEqual(stale.status_code, 200)
        self.assertTrue(stale.json().get("stale"))
        self.assertFalse(
            WhiteCartItem.objects.filter(
                cart__user=self.user,
                cart__status=WhiteCart.STATUS_ACTIVE,
            ).exists()
        )


class ColorfulHomePlacementTests(TestCase):
    def test_colorful_products_sit_below_categories_and_have_no_category_page(self):
        clothing = WhiteCategory.objects.create(name="ביגוד", slug="clothing", order=1)
        colorful = WhiteCategory.objects.create(
            name="צבעוני",
            slug="colorful",
            order=100,
            show_products_on_homepage=True,
        )
        WhiteSubcategory.objects.create(
            name="סט פיג'מה",
            slug="pajama-set",
            category=colorful,
        )

        home = self.client.get(reverse("white_catalog:home"))
        self.assertEqual(home.status_code, 200)
        html = home.content.decode()
        self.assertLess(html.find("ביגוד"), html.find("pajama-set"))
        nav_ids = list(home.context["all_categories"].values_list("id", flat=True))
        self.assertIn(clothing.id, nav_ids)
        self.assertNotIn(colorful.id, nav_ids)
        slugs = [product.slug for product in home.context["bottom_subcategories"]]
        self.assertIn("pajama-set", slugs)

        category_page = self.client.get(
            reverse("white_catalog:category_detail", kwargs={"category_slug": colorful.slug})
        )
        self.assertRedirects(category_page, reverse("white_catalog:home"))

        product = self.client.get(
            reverse(
                "white_catalog:subcategory_detail",
                kwargs={"category_slug": colorful.slug, "subcategory_slug": "pajama-set"},
            )
        )
        self.assertEqual(product.status_code, 200)
        self.assertNotContains(product, "צבעוני")


class ColorSizePerVariantTests(TestCase):
    def setUp(self):
        self.category = WhiteCategory.objects.create(name="צבעוני בדיקה", slug="colorful-sizes", order=1)
        self.product = WhiteSubcategory.objects.create(
            category=self.category,
            name="סט פיג'מה",
            slug="pajama-sizes",
            is_orderable=True,
            has_color_variants=True,
            unit_price=Decimal("23.00"),
        )
        self.color = WhiteColor.objects.create(name="ורוד בדיקה", hex_color="#ffc0cb")
        self.s1, _created = WhiteSizeType.objects.get_or_create(name="18-24", defaults={"order": 1})
        self.s2 = WhiteSizeType.objects.create(name="12-18-בדיקה", order=2)
        self.row1 = WhiteColorVariant.objects.create(
            product=self.product,
            color=self.color,
            size=self.s1,
            barcode="111222333",
            unit_price=Decimal("23.00"),
            order=1,
        )
        self.row2 = WhiteColorVariant.objects.create(
            product=self.product,
            color=self.color,
            size=self.s2,
            unit_price=Decimal("23.00"),
            order=2,
        )
        self.user = WhiteCatalogUser.objects.create(
            company_name="חנות",
            contact_name="בדיקה",
            contact_phone="0500000000",
            username="size-shop",
            password_hash="x",
        )
        session = self.client.session
        session["white_catalog_user_id"] = self.user.id
        session["white_catalog_username"] = self.user.username
        session.save()

    def test_product_page_asks_for_one_quantity_covering_every_size(self):
        response = self.client.get(
            reverse(
                "white_catalog:subcategory_detail",
                kwargs={"category_slug": self.category.slug, "subcategory_slug": self.product.slug},
            )
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "18-24")
        self.assertContains(response, "12-18-בדיקה")
        self.assertContains(response, f'name="color_set_qty_{self.color.id}"')
        self.assertNotContains(response, f'name="color_qty_{self.row1.id}"')
        self.assertContains(response, "111222333")
        self.assertContains(response, "הכמות זהה בכל המידות")

    def test_one_quantity_is_copied_onto_every_size(self):
        added = self.client.post(
            reverse("white_catalog:cart_add"),
            {
                "format": "json",
                "product_id": str(self.product.id),
                f"color_set_qty_{self.color.id}": "2",
            },
        )
        self.assertEqual(added.status_code, 200)
        self.assertEqual(added.json()["added"], 1)
        lines = list(WhiteCartItem.objects.filter(cart__user=self.user).order_by("color_variant__order"))
        self.assertEqual([line.quantity for line in lines], [2, 2])
        self.assertEqual(
            [line.color_variant.size.name for line in lines],
            ["18-24", "12-18-בדיקה"],
        )

        page = self.client.get(
            reverse(
                "white_catalog:subcategory_detail",
                kwargs={"category_slug": self.category.slug, "subcategory_slug": self.product.slug},
            )
        )
        self.assertContains(page, "2 מכל מידה = 4 סטים")

    def test_each_size_is_its_own_cart_and_order_line(self):
        added = self.client.post(
            reverse("white_catalog:cart_add"),
            {
                "format": "json",
                f"color_qty_{self.row1.id}": "2",
                f"color_qty_{self.row2.id}": "3",
            },
        )
        self.assertEqual(added.status_code, 200)
        self.assertEqual(added.json()["added"], 2)
        self.assertEqual(WhiteCartItem.objects.filter(cart__user=self.user).count(), 2)
        item = WhiteCartItem.objects.select_related("color_variant__color", "color_variant__size").get(
            color_variant=self.row1
        )
        self.assertEqual(item.display_size_name, "18-24")
        self.assertEqual(item.display_color_and_size, "ורוד בדיקה · 18-24")

        checkout = self.client.post(reverse("white_catalog:checkout"), {"notes": ""})
        self.assertEqual(checkout.status_code, 302)
        saved = WhiteOrderItem.objects.get(order__user=self.user, color_variant=self.row1)
        self.assertEqual(saved.color_name, "ורוד בדיקה")
        self.assertEqual(saved.size_name, "ורוד בדיקה · 18-24")

    def test_a_color_without_a_size_stays_a_single_row(self):
        other = WhiteColor.objects.create(name="לבן בדיקה")
        WhiteColorVariant.objects.create(product=self.product, color=other, unit_price=Decimal("10"))
        with self.assertRaises(IntegrityError):
            WhiteColorVariant.objects.create(product=self.product, color=other)

    def test_split_keeps_the_color_price_and_the_barcode_on_the_first_size(self):
        import importlib.util
        from pathlib import Path

        from django.apps import apps

        migration_path = Path(__file__).resolve().parent / "migrations" / "0032_color_variant_size.py"
        spec = importlib.util.spec_from_file_location("color_variant_size_migration", migration_path)
        migration = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(migration)

        product = WhiteSubcategory.objects.create(
            category=self.category,
            name="מניפה",
            slug="fan-split",
            is_orderable=True,
            has_color_variants=True,
        )
        color = WhiteColor.objects.create(name="כחול בדיקה")
        original = WhiteColorVariant.objects.create(
            product=product,
            color=color,
            barcode="999888777",
            unit_price=Decimal("23.00"),
        )
        WhiteSubcategoryPrice.objects.create(
            subcategory=product, size_name="24-30", price=Decimal("10"), order=1
        )
        WhiteSubcategoryPrice.objects.create(
            subcategory=product, size_name="30-36", price=Decimal("10"), order=2
        )
        migration.split_colors_into_sizes(apps, None)
        rows = list(
            WhiteColorVariant.objects.filter(product=product, color=color).select_related("size").order_by("order", "id")
        )
        self.assertEqual([row.size.name for row in rows], ["24-30", "30-36"])
        self.assertFalse(WhiteColorVariant.objects.filter(pk=original.pk).exists())
        self.assertEqual(rows[0].barcode, "999888777")
        self.assertIsNone(rows[1].barcode)
        self.assertEqual(rows[0].unit_price, Decimal("23.00"))
        self.assertEqual(rows[1].unit_price, Decimal("23.00"))


class ColorFilledOntoProductSizesTests(TestCase):
    def setUp(self):
        self.category = WhiteCategory.objects.create(name="צבעוני מילוי", slug="colorful-fill", order=1)
        self.product = WhiteSubcategory.objects.create(
            category=self.category,
            name="סט פיג'מה",
            slug="pajama-fill",
            is_orderable=True,
            has_color_variants=True,
            unit_price=Decimal("23.00"),
        )
        self.sizes = [
            WhiteSizeType.objects.get_or_create(name=name, defaults={"order": index})[0]
            for index, name in enumerate(("12-18", "18-24", "24-30"), start=1)
        ]
        for size in self.sizes:
            WhiteSubcategoryPrice.objects.create(
                subcategory=self.product,
                size_name=size.name,
                price=Decimal("23.00"),
                order=size.order,
            )
        self.existing = WhiteColor.objects.create(name="ורוד קיים")
        for size in self.sizes:
            WhiteColorVariant.objects.create(
                product=self.product,
                color=self.existing,
                size=size,
                unit_price=Decimal("23.00"),
                order=size.order,
            )

    def test_one_new_color_is_copied_onto_every_product_size(self):
        from white_catalog.models import fill_missing_color_sizes

        color = WhiteColor.objects.create(name="מנומר חום בדיקה")
        entered = WhiteColorVariant.objects.create(
            product=self.product,
            color=color,
            barcode="555666777",
            unit_price=Decimal("23.00"),
        )
        fill_missing_color_sizes(self.product, [color.id])
        rows = list(
            WhiteColorVariant.objects.filter(product=self.product, color=color)
            .select_related("size")
            .order_by("size__order")
        )
        self.assertEqual([row.size.name for row in rows], ["12-18", "18-24", "24-30"])
        self.assertEqual(rows[0].pk, entered.pk)
        self.assertEqual(rows[0].barcode, "555666777")
        self.assertIsNone(rows[1].barcode)
        self.assertIsNone(rows[2].barcode)
        self.assertEqual({row.unit_price for row in rows}, {Decimal("23.00")})

    def test_a_color_without_product_sizes_stays_one_row(self):
        from white_catalog.models import fill_missing_color_sizes

        plain = WhiteSubcategory.objects.create(
            category=self.category,
            name="מגבת",
            slug="plain-towel",
            is_orderable=True,
            has_color_variants=True,
        )
        color = WhiteColor.objects.create(name="לבן מילוי")
        WhiteColorVariant.objects.create(product=plain, color=color)
        fill_missing_color_sizes(plain, [color.id])
        self.assertEqual(WhiteColorVariant.objects.filter(product=plain, color=color).count(), 1)
        self.assertFalse(WhiteColorVariant.objects.filter(product=plain, color=color, size__isnull=False).exists())

    def test_editing_an_existing_size_does_not_ask_to_fill_again(self):
        from white_catalog.admin import colors_to_expand

        class FakeForm:
            def __init__(self, color, size, pk=None, delete=False):
                self.cleaned_data = {"color": color, "size": size, "DELETE": delete}
                self.instance = type("Instance", (), {"pk": pk})()

        existing = FakeForm(self.existing, self.sizes[0], pk=10)
        removed = FakeForm(self.existing, self.sizes[2], pk=12, delete=True)
        fresh = WhiteColor.objects.create(name="תכלת מילוי")
        added = FakeForm(fresh, None, pk=None)
        self.assertEqual(colors_to_expand([existing, removed, added]), [fresh.id])


class ColorAutocompleteSwatchTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_superuser(
            username="swatch-admin",
            email="swatch@example.com",
            password="secret",
        )
        self.client.force_login(self.admin)

    def test_color_search_returns_the_swatch_and_the_short_name(self):
        from django.core.files.base import ContentFile

        photo = WhiteColor.objects.create(name="מנומר תפריט", notes="בד אינטרלוק אריה מנומר")
        photo.swatch_image.save("swatch.gif", ContentFile(
            b"GIF89a\x01\x00\x01\x00\x80\x00\x00\x00\x00\x00\xff\xff\xff!"
            b"\xf9\x04\x01\x00\x00\x00\x00,\x00\x00\x00\x00\x01\x00\x01\x00\x00\x02\x02D\x01\x00;"
        ), save=True)
        plain = WhiteColor.objects.create(name="אפור תפריט", notes="הערה ארוכה", hex_color="#dedede")
        WhiteSizeType.objects.create(name="12-18-תפריט")

        response = self.client.get(
            reverse("admin:autocomplete"),
            {
                "term": "תפריט",
                "app_label": "white_catalog",
                "model_name": "whitecolorvariant",
                "field_name": "color",
            },
        )
        self.assertEqual(response.status_code, 200)
        by_id = {row["id"]: row for row in response.json()["results"]}
        self.assertEqual(by_id[str(photo.pk)]["text"], "מנומר תפריט")
        self.assertNotIn("אינטרלוק", by_id[str(photo.pk)]["text"])
        self.assertIn("swatch.gif", by_id[str(photo.pk)]["swatch"])
        self.assertEqual(by_id[str(plain.pk)]["text"], "אפור תפריט")
        self.assertEqual(by_id[str(plain.pk)]["hex"], "#dedede")
        self.assertNotIn("swatch", by_id[str(plain.pk)])

        sizes = self.client.get(
            reverse("admin:autocomplete"),
            {
                "term": "12-18-תפריט",
                "app_label": "white_catalog",
                "model_name": "whitecolorvariant",
                "field_name": "size",
            },
        )
        self.assertEqual(sizes.status_code, 200)
        size_row = sizes.json()["results"][0]
        self.assertEqual(size_row["text"], "12-18-תפריט")
        self.assertNotIn("swatch", size_row)


